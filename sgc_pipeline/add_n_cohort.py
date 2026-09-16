#!/usr/bin/env python3
"""
Add an n_cohort column to the Clump_index_variants sheet.

For each row (GWAS, index_variant), look up the variant's position in the
matching meta sumstat file <GWAS>.tsv.gz in --gwas-dir, read its n_cohort
value, and write it into a new column in the sheet. Position-only matching
(chr:pos), consistent with the rest of the pipeline.

Each sumstat file is scanned once for all the variants that need it.

Usage:
    python3 add_n_cohort.py \\
        --xlsx SGC_Meta_Analysis_Results.xlsx \\
        --gwas-dir . \\
        --out SGC_Meta_Analysis_Results.xlsx \\
        --threads 8
"""

import argparse
import gzip
import io
import os
import re
import sys
from collections import defaultdict
from concurrent.futures import ProcessPoolExecutor
from copy import copy

SHEET = "Clump_index_variants"
CHR_ALIASES = {"CHR", "CHROMOSOME", "CHROM", "#CHROM"}
POS_ALIASES = {"POS", "POSITION", "GENPOS", "BP", "BASE_PAIR_LOCATION"}


def norm_chr(c):
    c = re.sub(r"^chr", "", str(c).strip(), flags=re.I).upper()
    return {"23": "X", "24": "Y", "25": "X", "26": "MT", "M": "MT"}.get(c, c)


def open_text(path):
    if path.endswith(".gz"):
        return io.TextIOWrapper(gzip.open(path, "rb"), encoding="utf-8",
                                errors="replace")
    return open(path, encoding="utf-8", errors="replace")


def sniff_delim(line):
    if "\t" in line:
        return "\t"
    if "," in line:
        return ","
    return None


def variant_pos(variant):
    """chr:pos[:a1:a2] -> (chr, pos) ; position-only."""
    p = str(variant).split(":")
    if len(p) < 2:
        return None
    try:
        return (norm_chr(p[0]), int(p[1]))
    except ValueError:
        return None


def scan_file(job):
    """Read one sumstat file once; return {(chr,pos): n_cohort} for wanted keys."""
    path, want, ncol_name = job
    fn = os.path.basename(path)
    out = {}
    if not os.path.exists(path):
        return (fn, None, "file not found")
    try:
        with open_text(path) as fh:
            first = fh.readline()
            delim = sniff_delim(first)
            header = (first.rstrip("\n").split(delim) if delim
                      else first.split())
            idx = {h.strip().upper(): i for i, h in enumerate(header)}
            ic = next((idx[a] for a in CHR_ALIASES if a in idx), None)
            ip = next((idx[a] for a in POS_ALIASES if a in idx), None)
            inc = idx.get(ncol_name.upper())
            if ic is None or ip is None:
                return (fn, None, f"no chr/pos column; header={header[:6]}")
            if inc is None:
                return (fn, None, f"no {ncol_name!r} column; header={header[:8]}")
            need = max(ic, ip, inc)
            for line in fh:
                f = line.rstrip("\n").split(delim) if delim else line.split()
                if len(f) <= need:
                    continue
                try:
                    key = (norm_chr(f[ic]), int(f[ip]))
                except ValueError:
                    continue
                if key in want and key not in out:
                    out[key] = f[inc].strip()
    except OSError as e:
        return (fn, None, f"read error: {e}")
    return (fn, out, None)


def main():
    ap = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--xlsx", required=True)
    ap.add_argument("--sheet", default=SHEET)
    ap.add_argument("--gwas-dir", default=".")
    ap.add_argument("--suffix", default=".tsv.gz")
    ap.add_argument("--ncol", default="n_cohort",
                    help="name of the n_cohort column in the sumstat files")
    ap.add_argument("--out", required=True)
    ap.add_argument("--threads", type=int, default=0,
                    help="parallel processes (0 = all cores)")
    args = ap.parse_args()

    from openpyxl import load_workbook
    wb = load_workbook(args.xlsx)
    if args.sheet not in wb.sheetnames:
        sys.exit(f"ERROR: sheet {args.sheet!r} not found; have {wb.sheetnames}")
    ws = wb[args.sheet]
    hdr = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    H = {h: i + 1 for i, h in enumerate(hdr) if h is not None}
    for need in ("GWAS", "index_variant"):
        if need not in H:
            sys.exit(f"ERROR: column {need!r} not in sheet; have {hdr}")
    c_gwas, c_var = H["GWAS"], H["index_variant"]

    # collect, per GWAS file, the set of positions we need
    per_gwas = defaultdict(set)       # gwas -> {(chr,pos)}
    row_key = {}                      # row -> (gwas, chr, pos)
    for r in range(2, ws.max_row + 1):
        gwas = ws.cell(r, c_gwas).value
        var = ws.cell(r, c_var).value
        if gwas is None or var is None:
            continue
        pk = variant_pos(var)
        if pk is None:
            continue
        gwas = str(gwas).strip()
        per_gwas[gwas].add(pk)
        row_key[r] = (gwas, pk)
    sys.stderr.write(f"{len(row_key)} variant rows across {len(per_gwas)} GWAS\n")

    jobs = []
    for gwas, keys in per_gwas.items():
        path = os.path.join(args.gwas_dir, f"{gwas}{args.suffix}")
        jobs.append((path, keys, args.ncol))

    n_threads = args.threads if args.threads and args.threads > 0 \
        else (os.cpu_count() or 1)
    n_threads = min(n_threads, len(jobs))
    sys.stderr.write(f"Scanning {len(jobs)} sumstat files on {n_threads} "
                     f"process(es)\n")

    # gwas -> {(chr,pos): n_cohort}
    results = {}
    missing_files = []
    done = 0

    def handle(gwas, res):
        nonlocal done
        fn, out, err = res
        done += 1
        if err:
            sys.stderr.write(f"WARN [{fn}] {err}\n")
            if "not found" in err:
                missing_files.append(gwas)
        else:
            results[gwas] = out
        if done % 25 == 0 or done == len(jobs):
            sys.stderr.write(f"  scanned {done}/{len(jobs)}\n")

    gwas_order = list(per_gwas.keys())
    if n_threads == 1:
        for gwas, job in zip(gwas_order, jobs):
            handle(gwas, scan_file(job))
    else:
        with ProcessPoolExecutor(max_workers=n_threads) as ex:
            for gwas, res in zip(gwas_order, ex.map(scan_file, jobs)):
                handle(gwas, res)

    # insert the new column right after index_variant
    new_col = c_var + 1
    ws.insert_cols(new_col)
    hstyle, bstyle = ws.cell(1, 1), ws.cell(2, 1)
    hc = ws.cell(1, new_col, "n_cohort")
    hc.font = copy(hstyle.font); hc.fill = copy(hstyle.fill)
    hc.alignment = copy(hstyle.alignment); hc.border = copy(hstyle.border)

    # inserting shifted columns >= new_col right by one, but row_key is keyed by
    # row number (unchanged) and holds the values we already read, so we just
    # fill the new column from row_key + results.
    filled = blank = 0
    for r in range(2, ws.max_row + 1):
        cell = ws.cell(r, new_col)
        cell.font = copy(bstyle.font)
        rk = row_key.get(r)
        if rk is None:
            blank += 1
            continue
        gwas, pk = rk
        val = results.get(gwas, {}).get(pk)
        if val is None or val == "":
            blank += 1
        else:
            # numeric if possible
            try:
                cell.value = int(val)
            except ValueError:
                try:
                    cell.value = float(val)
                except ValueError:
                    cell.value = val
            filled += 1

    wb.save(args.out)
    sys.stderr.write(
        f"\nn_cohort added at column {new_col}.\n"
        f"filled: {filled} | blank: {blank}\n")
    if missing_files:
        uniq = sorted(set(missing_files))
        sys.stderr.write(f"{len(uniq)} GWAS had no sumstat file: "
                         + ", ".join(uniq[:10])
                         + (" ..." if len(uniq) > 10 else "") + "\n")
    sys.stderr.write(f"Wrote {args.out}\n")


if __name__ == "__main__":
    main()
